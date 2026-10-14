"""Benchmark records, loaders for open datasets, and campaign-aware splits.

Every record is re-redacted on load (victims' OTPs, cards, Aadhaar, PAN), keeps its source and
licence, and starts unverified: imported labels are noisy (a genuine vaccination reminder labelled
as a scam, for instance), so only records a person has checked count as benchmark ground truth.

Splits are made by template, not by message. Scam operations send thousands of near-copies; if
copies of one template sit in both training and test data, a detector is rewarded for memorising
it. Near-copies (and messages whose translations are near-copies) are grouped first, and whole
groups are assigned to a split.
"""
import csv
import hashlib
import json
import random
import re
from collections import defaultdict
from dataclasses import asdict, dataclass, field

from .campaigns import _DSU, shingles
from .indicators import redact
from .lsh import near_copies
from .taxonomy import IMC25_TYPES, language

IMC25_URL = ("https://raw.githubusercontent.com/reportsmishing/Smishing-Dataset-IMC25/main/"
             "dataset/final_dataset_output.csv")
IMC25_SHA256 = "1bbd1e9e82c3ea023112207b80da268a5c4a07d2353c2b0898360ab037fa9a64"   # main branch, fetched 2026-10-14
IMC25_CITATION = ("Agarwal, Papasavva, Suarez-Tangil, Vasek. Fishing for Smishing: Understanding SMS "
                  "Phishing Infrastructure and Strategies by Mining Public User Reports. IMC 2025. "
                  "doi:10.1145/3730567.3764431. CC BY 4.0.")


@dataclass
class Record:
    id: str
    text: str
    lang: str
    region: str
    label: str                       # scam | spam | legit | unknown
    scam_type: str = None
    lures: list = field(default_factory=list)
    translation: str = ""           # English, when the source provides one
    variety: str = "native"          # native | romanized | code-mixed (how the language is written)
    country: str = ""                # ISO 3166 alpha-2, when known
    source: str = ""
    license: str = ""
    verified: bool = False
    group: int = -1                  # template group, for splits

    def to_json(self):
        return json.dumps(asdict(self), ensure_ascii=False)


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def load_imc25(path):
    """Records from the IMC 2025 smishing dataset (CC BY 4.0). Every report in it is a message
    users reported as smishing; its 'spam' type marks plain marketing and becomes label 'spam'."""
    csv.field_size_limit(1 << 30)
    out = []
    with open(path, encoding="utf-8", errors="replace", newline="") as f:
        for n, r in enumerate(csv.DictReader(f)):
            text = (r.get("text") or "").strip()
            if not text:
                continue
            code, region = language(r.get("language"))
            t = IMC25_TYPES.get((r.get("scam_type") or "").strip().lower())
            label = "spam" if t == "spam" else "scam"
            lures = [x.strip() for x in (r.get("lure_principles") or "").strip("[]").replace("'", "").split(",")
                     if x.strip()]
            out.append(Record(id=f"imc25-{n}", text=redact(text), lang=code, region=region, label=label,
                              scam_type=None if t == "spam" else t, lures=lures,
                              translation=redact((r.get("translation") or "").strip()),
                              source="imc25", license="CC-BY-4.0"))
    return out


TAG = re.compile(r"<[A-Z_]+>|\[(?:OTP|CARD|AADHAAR|PAN)\]")
MIN_CONTENT = 20     # characters of real text needed to compare a message with others


def content(text):
    """Text without anonymisation placeholders; '' when too little real text is left to compare
    (messages that are mostly placeholders would otherwise chain unrelated scams together)."""
    t = " ".join(TAG.sub(" ", text or "").split())
    return t if len(t) >= MIN_CONTENT else ""


def english_view(rec):
    return rec.translation or (rec.text if rec.lang == "en" else "")


def group_templates(records, threshold=0.6):
    """Assign rec.group: near-copies in the original text, or in the English translation, share a
    group. Returns the number of groups."""
    n = len(records)
    dsu = _DSU(n)
    for view in (lambda r: content(r.text), lambda r: content(english_view(r))):
        sets = [shingles(v) if (v := view(r)) else set() for r in records]
        for i, j, _ in near_copies(sets, threshold):
            dsu.union(i, j)
    roots = {}
    for i, r in enumerate(records):
        r.group = roots.setdefault(dsu.find(i), len(roots))
    return len(roots)


def split(records, fractions=(("train", 0.7), ("dev", 0.1), ("test", 0.2)), seed=13):
    """{split name: [records]} with every template group wholly inside one split."""
    groups = defaultdict(list)
    for r in records:
        groups[r.group].append(r)
    keys = sorted(groups)
    random.Random(seed).shuffle(keys)
    total, out, acc = len(records), {name: [] for name, _ in fractions}, 0
    bounds, run = [], 0.0
    for name, frac in fractions:
        run += frac
        bounds.append((name, run * total))
    for k in keys:
        name = next((nm for nm, b in bounds if acc < b), bounds[-1][0])
        out[name].extend(groups[k])
        acc += len(groups[k])
    return out


def load_bengali_smishing(directory):
    """Bengali SMS Smishing Dataset (MIT): Bengali script, Banglish (Bengali in Latin letters),
    code-mixed and English. smish -> scam, promo -> spam, normal -> legit."""
    import pyarrow.parquet as pq
    labels = {"smish": "scam", "promo": "spam", "normal": "legit"}
    variety = {"Bengali": ("bn", "native"), "Banglish": ("bn", "romanized"),
               "CodeMix": ("bn", "code-mixed"), "English": ("en", "native")}
    out = []
    for split_name in ("train", "validation", "test"):
        rows = pq.read_table(f"{directory}/{split_name}.parquet").to_pylist()
        for n, r in enumerate(rows):
            lang, var = variety.get(r["source"], ("und", "native"))
            out.append(Record(id=f"bnsmish-{split_name}-{n}", text=redact(r["text"]), lang=lang,
                              region="south_asia" if lang == "bn" else "global",
                              label=labels.get(r["label"], "unknown"), variety=var,
                              source="bengali_smishing", license="MIT"))
    return out


def load_moz_smishing(path):
    """MOZ-Smishing (AfricaNLP 2025): Mozambican Portuguese, mobile money. Licensed
    creativeml-openrail-m and containing names: evaluated locally, never redistributed."""
    out = []
    with open(path, encoding="utf-8", errors="replace", newline="") as f:
        for r in csv.DictReader(f):
            out.append(Record(id=f"moz-{r['id']}", text=redact(r["text"]), lang="pt", region="africa",
                              label="scam" if r["label"].strip().lower() == "smishing" else "legit",
                              scam_type="bank_payment" if r["label"].strip().lower() == "smishing" else None,
                              country="MZ", source="moz_smishing", license="creativeml-openrail-m"))
    return out


def load_singlish(path):
    """Translated Singlish SMS (CC BY 4.0): everyday messages from the NUS SMS Corpus with native
    English translations. All genuine, so they serve as hard negatives in Southeast Asia."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb["ALL_TRANSLATIONS"]
    out = []
    for n, row in enumerate(ws.iter_rows(values_only=True)):
        text, translation = (row[0] or "").strip(), (row[1] or "").strip()
        if not text:
            continue
        out.append(Record(id=f"singlish-{n}", text=redact(text), lang="en", region="southeast_asia",
                          label="legit", translation=redact(translation), variety="code-mixed",
                          country="SG", source="singlish_sms", license="CC-BY-4.0"))
    return out


def review_sample(records, per_label=100, seed=7, labels=("scam", "spam", "legit")):
    """Records for native speakers to check: for each language (and way of writing it) and each
    imported label, up to `per_label` records. One record per template group comes first, so large
    campaigns cannot crowd out the rest; only then are second copies taken from the largest groups."""
    rng = random.Random(seed)
    cells = defaultdict(lambda: defaultdict(list))
    for r in records:
        if r.label in labels and r.lang != "und":
            cells[(r.lang, r.variety, r.label)][r.group].append(r)
    out = []
    for key in sorted(cells):
        groups = list(cells[key].values())
        rng.shuffle(groups)
        for g in groups:
            rng.shuffle(g)
        groups.sort(key=len, reverse=True)
        picked, depth = [], 0
        while len(picked) < per_label and any(len(g) > depth for g in groups):
            batch = [g[depth] for g in groups if len(g) > depth]
            if depth == 0:
                rng.shuffle(batch)
            picked.extend(batch[:per_label - len(picked)])
            depth += 1
        out.extend(picked)
    return out
